from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.export.kigs_schema import load_export_profile
from app.services.export import KigsExportService
from test_phase5 import _phase5_session


def _sheet_rows(path: Path, name: str):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[name]
        headers = [cell.value for cell in next(sheet.iter_rows())]
        return headers, [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()


def test_kigs_export_structure_population_hash_and_versioning(client, sample_files):
    reconciliation_id = _phase5_session(client, sample_files)
    review = client.get(f"/api/reconciliations/{reconciliation_id}/final-review")
    assert review.status_code == 200 and review.json()["validation"]["valid"] is True
    assert review.json()["exact_matches"] == 520
    assert review.json()["tolerance_matches"] == 140
    assert review.json()["near_matches"] == 100
    assert review.json()["resolved_records"] == 760
    assert review.json()["active_rules"] == []

    first = client.post(f"/api/reconciliations/{reconciliation_id}/exports")
    assert first.status_code == 200, first.text
    record = first.json()
    path = Path(record["file_path"])
    assert path.is_file() and sha256(path.read_bytes()).hexdigest() == record["sha256"]
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["KIGS_Reconciliation", "Summary", "Unresolved_Exceptions", "Configuration", "Audit_Summary"]
    finally:
        workbook.close()
    headers, rows = _sheet_rows(path, "KIGS_Reconciliation")
    for confirmed in ("LocationGstin", "ReconciliationSection", "Reason", "ActionStatus", "Action",
                      "CPGstin", "PRGstin", "CPDocumentNumber", "PRDocumentNumber",
                      "TaxableDifference", "ReconciliationDateTime"):
        assert confirmed in headers
    assert len(rows) == record["row_count"] == 1130

    gst_only = next(row for row in rows if row["ReconciliationSection"] == "GST Only")
    assert gst_only["CPGstin"] and gst_only["CPDocumentNumber"] and gst_only["PRGstin"] is None
    assert gst_only["PRDocumentNumber"] is None and gst_only["ActionStatus"] == "Actions Not Taken"
    pr_only = next(row for row in rows if row["ReconciliationSection"] == "PR Only")
    assert pr_only["PRGstin"] and pr_only["PRDocumentNumber"] and pr_only["CPGstin"] is None
    mismatch = next(row for row in rows if row["ReconciliationSection"] == "Mismatched")
    assert mismatch["CPGstin"] and mismatch["PRGstin"]
    assert mismatch["TaxableDifference"] == mismatch["PRTaxableValue"] - mismatch["CPTaxableValue"]
    matched = next(row for row in rows if row["ReconciliationSection"] == "Matched")
    assert matched["CPGstin"] and matched["PRGstin"]
    assert matched["TaxableDifference"] == matched["PRTaxableValue"] - matched["CPTaxableValue"]

    summary_headers, summary_rows = _sheet_rows(path, "Summary")
    summary = {row["Metric"]: row["Value"] for row in summary_rows}
    assert summary["Government Record Count"] == 1000 and summary["Purchase Register Record Count"] == 1050
    assert summary["Exact Matches"] == 520 and summary["Near Matches"] == 100
    second = client.post(f"/api/reconciliations/{reconciliation_id}/exports").json()
    assert second["version"] == 2 and second["export_id"] != record["export_id"]
    assert len(client.get(f"/api/reconciliations/{reconciliation_id}/exports").json()) == 2


def test_human_selected_metadata_and_formula_text_safety(client, sample_files):
    reconciliation_id = _phase5_session(client, sample_files)
    tools = client.app.state.exception_tools
    ambiguity = tools.get_record(reconciliation_id, "GST-00841")
    selected_pr = ambiguity.best_candidate.purchase_register_record_id
    selected = client.post(f"/api/reconciliations/{reconciliation_id}/ambiguous/GST-00841/select", json={
        "action": "select", "purchase_register_record_id": selected_pr,
    })
    assert selected.status_code == 200 and selected.json()["resolved_records"] == 761
    export = client.post(f"/api/reconciliations/{reconciliation_id}/exports").json()
    _, rows = _sheet_rows(Path(export["file_path"]), "KIGS_Reconciliation")
    human = next(row for row in rows if row["ReferenceId"] == "GST-00841")
    assert human["ManualReconciliation"] == "Yes" and human["ReconciledBy"] == "POC user"
    assert human["CPDocumentNumber"] and human["PRDocumentNumber"]
    assert human["ReconciliationDateTime"] is not None
    assert KigsExportService._safe("=2+2") == "'=2+2"
    assert KigsExportService._safe("+cmd") == "'+cmd"
    assert KigsExportService._safe(-100) == -100


def test_template_header_order_unknown_fields_and_stale_history(client, sample_files, tmp_path):
    template = tmp_path / "header-only.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KIGS_Reconciliation"
    expected = ["Reason", "UnknownOfficialField", "ReconciliationSection", "ActionStatus", "Action",
                "CPGstin", "PRGstin", "CPDocumentNumber", "PRDocumentNumber"]
    sheet.append(expected)
    workbook.save(template)
    profile = load_export_profile(template)
    assert profile.field_order == expected and profile.schema_source == "TEMPLATE"
    unknown = next(item for item in profile.fields if item.target_name == "UnknownOfficialField")
    assert unknown.passthrough is True and unknown.confirmed is False

    reconciliation_id = _phase5_session(client, sample_files)
    generated = client.post(f"/api/reconciliations/{reconciliation_id}/exports")
    assert generated.status_code == 200
    assert client.post(f"/api/reconciliations/{reconciliation_id}/policy/propose").status_code == 200
    history = client.get(f"/api/reconciliations/{reconciliation_id}/exports").json()
    assert history[0]["stale"] is True and Path(history[0]["file_path"]).is_file()
    blocked = client.post(f"/api/reconciliations/{reconciliation_id}/exports")
    assert blocked.status_code == 409
    assert any(item["code"] == "policy_not_confirmed" for item in blocked.json()["detail"]["validation"]["issues"])


def test_upload_safety_and_bounded_results(client, sample_files):
    reconciliation_id = client.post("/api/reconciliations").json()["id"]
    with sample_files["government"].open("rb") as workbook:
        uploaded = client.post(
            f"/api/reconciliations/{reconciliation_id}/files/government",
            files={"file": ("../../outside.xlsx", workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert uploaded.status_code == 200
    assert uploaded.json()["original_filename"] == "outside.xlsx"
    stored = Path(uploaded.json()["stored_path"]).resolve()
    assert stored.parent.name == reconciliation_id

    rejected_extension = client.post(
        f"/api/reconciliations/{reconciliation_id}/files/purchase-register",
        files={"file": ("macro.xlsm", b"not a workbook", "application/octet-stream")},
    )
    assert rejected_extension.status_code == 415
    corrupt = client.post(
        f"/api/reconciliations/{reconciliation_id}/files/purchase-register",
        files={"file": ("corrupt.xlsx", b"not a workbook", "application/octet-stream")},
    )
    assert corrupt.status_code == 422

    completed_id = _phase5_session(client, sample_files)
    first_page = client.get(f"/api/reconciliations/{completed_id}/results?offset=0&limit=25")
    assert first_page.status_code == 200
    payload = first_page.json()
    assert len(payload["records"]) == 25
    assert payload["total_records"] > len(payload["records"])
    assert payload["offset"] == 0 and payload["limit"] == 25
    assert client.get(f"/api/reconciliations/{completed_id}/results?limit=1001").status_code == 422
