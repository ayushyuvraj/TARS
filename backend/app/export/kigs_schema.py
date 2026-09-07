from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.domain.models import (
    ExportFieldDefinition, ExportFieldGroup, ExportProfile, ExportSchemaSource,
)


RECONCILIATION_FIELDS = [
    "LocationGstin", "LocationName", "ReconciliationSection", "SuggReconciliationSection",
    "Reason", "ActionStatus", "Action", "ValueDifference", "TaxableDifference",
    "IGSTDifference", "CGSTDifference", "SGSTDifference", "CessDifference",
]
CP_FIELDS = [
    "CPGstin", "CPTradeName", "CPLegalName", "CPDocumentNumber", "CPDocumentDate",
    "CPDocumentType", "CPValue", "CPPOS", "CPReverseCharge", "CPTaxableValue",
    "CPRate", "CPIgstAmount", "CPCgstAmount", "CPSgstAmount", "CPCessAmount",
]
PR_FIELDS = [
    "PRGstin", "PRTradeName", "PRLegalName", "PRDocumentNumber", "PRDocumentDate",
    "PRDocumentType", "PRValue", "PRPOS", "PRReverseCharge", "PRTaxableValue",
    "PRRate", "PRIgstAmount", "PRCgstAmount", "PRSgstAmount", "PRCessAmount",
]
PROCESS_FIELDS = [
    "ReconciliationPercentage", "ReconciledBy", "ManualReconciliation",
    "ReconciliationDateTime", "Remarks", "IMSAction", "IMSRemarks", "Gstr1GrcScore",
    "Gstr3bGrcScore", "Notes", "ReduceITC", "ReferenceId", "CPIsReversal",
    "Gstr3bNonPayment",
]
CONFIGURED_FIELDS = RECONCILIATION_FIELDS + CP_FIELDS + PR_FIELDS + PROCESS_FIELDS
REQUIRED_FIELDS = [
    "ReconciliationSection", "Reason", "ActionStatus", "Action",
    "CPGstin", "PRGstin", "CPDocumentNumber", "PRDocumentNumber",
]

STATUS_MAPPING = {
    "EXACT_MATCHED": "Matched",
    "TOLERANCE_MATCHED": "Matched",
    "NEAR_MATCHED": "Matched",
    "HUMAN_SELECTED": "Matched",
    "MATERIAL_MISMATCH": "Mismatched",
    "AMBIGUOUS": "Unreconciled",
    "GST_ONLY": "GST Only",
    "PR_ONLY": "PR Only",
    "UNRESOLVED": "Unreconciled",
}
ACTION_MAPPING = {"default_status": "Actions Not Taken", "default_action": "No Action"}

NUMBER_FIELDS = {
    "ValueDifference", "TaxableDifference", "IGSTDifference", "CGSTDifference",
    "SGSTDifference", "CessDifference", "CPValue", "CPTaxableValue", "CPRate",
    "CPIgstAmount", "CPCgstAmount", "CPSgstAmount", "CPCessAmount", "PRValue",
    "PRTaxableValue", "PRRate", "PRIgstAmount", "PRCgstAmount", "PRSgstAmount",
    "PRCessAmount", "ReconciliationPercentage", "Gstr1GrcScore", "Gstr3bGrcScore",
}
DATE_FIELDS = {"CPDocumentDate", "PRDocumentDate"}
DATETIME_FIELDS = {"ReconciliationDateTime"}


def _group(name: str) -> ExportFieldGroup:
    if name in RECONCILIATION_FIELDS:
        return ExportFieldGroup.RECONCILIATION_DERIVED
    if name in CP_FIELDS:
        return ExportFieldGroup.CP_GOVERNMENT
    if name in PR_FIELDS:
        return ExportFieldGroup.PR_PURCHASE_REGISTER
    return ExportFieldGroup.PROCESS_METADATA


def _field(name: str, known: bool = True) -> ExportFieldDefinition:
    datatype = "number" if name in NUMBER_FIELDS else (
        "date" if name in DATE_FIELDS else "datetime" if name in DATETIME_FIELDS else "text"
    )
    return ExportFieldDefinition(
        target_name=name, group=_group(name), datatype=datatype,
        required=name in REQUIRED_FIELDS, confirmed=known, passthrough=not known,
    )


def _template_headers(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        worksheet = workbook["KIGS_Reconciliation"] if "KIGS_Reconciliation" in workbook.sheetnames else workbook.worksheets[0]
        for row in worksheet.iter_rows(min_row=1, max_row=30, values_only=True):
            headers = [str(value).strip() if value is not None else "" for value in row]
            while headers and not headers[-1]:
                headers.pop()
            nonblank = [value for value in headers if value]
            if len(nonblank) >= 2 and len(nonblank) == len(set(nonblank)):
                if any(value in REQUIRED_FIELDS for value in nonblank):
                    return headers
    finally:
        workbook.close()
    raise ValueError("The KIGS template does not contain a recognizable unique header row")


def load_export_profile(template_path: Path | None = None) -> ExportProfile:
    source = ExportSchemaSource.CONFIGURED_POC
    headers = list(CONFIGURED_FIELDS)
    if template_path and template_path.is_file():
        headers = _template_headers(template_path)
        source = ExportSchemaSource.TEMPLATE
    fields = [_field(name, name in CONFIGURED_FIELDS) for name in headers]
    return ExportProfile(
        schema_source=source, field_order=headers, fields=fields,
        required_fields=REQUIRED_FIELDS, status_mapping=STATUS_MAPPING,
        action_mapping=ACTION_MAPPING,
    )
