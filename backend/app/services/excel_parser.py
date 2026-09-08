from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

from app.domain.models import CanonicalDataType, ColumnProfile, DatasetProfile, DatasetRole, RoleDetectionResult


class ExcelParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDataset:
    dataframe: pd.DataFrame
    profile: DatasetProfile


class ExcelParser:
    GSTIN_PATTERN = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][0-9A-Z]Z[0-9A-Z]$")

    GOV_MARKERS = {
        "gstinofsupplier", "suppliergstin", "ctin", "tradelegalname", "legalname",
        "gstr15filingdate", "gstr15filingperiod", "invoicetype", "placeofsupply",
        "reversecharge", "integratedtax", "centraltax", "stateuttax", "taxablevalue",
        "invoicevalue", "invoicenumber", "invoicedate", "itcavailability", "itc",
    }
    PR_MARKERS = {
        "vendorcode", "vendorname", "partyname", "ponumber", "vouchernumber",
        "voucherno", "postingdate", "billnumber", "billno", "purchaseaccount",
        "businessunit", "costcenter", "ledgername", "internalref", "entrydate",
        "documentno", "companycode", "purchaseorder",
    }

    def __init__(self) -> None:
        self._cache: dict[tuple[str, str, float], ParsedDataset] = {}

    def clear_cache(self) -> None:
        self._cache.clear()

    def detect_header_row(self, path: Path, role: DatasetRole = DatasetRole.GOVERNMENT) -> tuple[str, int]:

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise ExcelParseError("The uploaded file is not a readable Excel workbook") from exc
        try:
            worksheet = workbook.worksheets[0]
            candidates: list[tuple[int, int]] = []
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1
            ):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                unique_values = set(values)
                if len(values) >= 2 and len(unique_values) == len(values):
                    candidates.append((len(values), row_number))
            if candidates:
                _, header_row = max(candidates, key=lambda candidate: (candidate[0], -candidate[1]))
                return worksheet.title, header_row
        finally:
            workbook.close()
        raise ExcelParseError("Could not locate a tabular header in the first 30 rows")

    def _header_scores(self, path: Path) -> tuple[int, int]:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                worksheet = workbook.worksheets[0]
                row_iterator = worksheet.iter_rows(min_row=1, max_row=30, values_only=True)
                candidates: list[tuple[int, int, list[str]]] = []
                for row_number, row in enumerate(row_iterator, start=1):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    unique_values = set(values)
                    if len(values) >= 2 and len(unique_values) == len(values):
                        candidates.append((len(values), row_number, values))
                if not candidates:
                    return 0, 0
                _, _, header_vals = max(candidates, key=lambda candidate: (candidate[0], -candidate[1]))
                headers = [re.sub(r"[^a-z0-9]", "", str(c).lower()) for c in header_vals]
            finally:
                workbook.close()
            fname = path.name.lower()
            gov_score = sum(1 for h in headers if any(m in h for m in self.GOV_MARKERS))
            pr_score = sum(1 for h in headers if any(m in h for m in self.PR_MARKERS))
            if any(k in fname for k in ("2b", "gstr", "govt", "government")):
                gov_score += 3
            if any(k in fname for k in ("pr", "purchase", "register", "books")):
                pr_score += 3
            return gov_score, pr_score
        except Exception:
            return 0, 0

    def _read_frame(self, path: Path, header_row: int) -> pd.DataFrame:
        try:
            return pd.read_excel(path, sheet_name=0, header=header_row - 1, engine="calamine")
        except Exception:
            pass
        try:
            return pd.read_excel(path, sheet_name=0, header=header_row - 1)
        except Exception:
            pass
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                worksheet = workbook.worksheets[0]
                data = list(worksheet.values)
                if not data or len(data) < header_row:
                    raise ExcelParseError("The worksheet header row is empty")
                header_values = data[header_row - 1]
                columns = [str(col).strip() if col is not None else "" for col in header_values]
                rows_data = data[header_row:]
                return pd.DataFrame(rows_data, columns=columns)
            finally:
                workbook.close()
        except ExcelParseError:
            raise
        except Exception as exc:
            raise ExcelParseError("Failed to parse the Excel worksheet") from exc

    def parse(self, path: Path, role: DatasetRole) -> ParsedDataset:
        resolved_path = str(path.resolve())
        mtime = path.stat().st_mtime if path.exists() else 0.0
        cache_key = (resolved_path, role.value, mtime)
        if cache_key in self._cache:
            return self._cache[cache_key]

        sheet_name, header_row = self.detect_header_row(path, role)
        frame = self._read_frame(path, header_row)
        frame = frame.dropna(how="all").reset_index(drop=True)
        frame.columns = [str(column).strip() for column in frame.columns]
        if len(frame.columns) < 2:
            raise ExcelParseError("The worksheet must contain at least two columns")
        if len(set(frame.columns)) != len(frame.columns):
            raise ExcelParseError("Duplicate source column names are not supported")

        row_count = len(frame)
        nunique_dict = frame.nunique(dropna=True).to_dict()
        null_counts = frame.isna().sum().to_dict()

        column_profiles = [
            self._profile_column(
                column, frame[column], row_count, nunique_dict.get(column, 0), null_counts.get(column, 0)
            )
            for column in frame.columns
        ]
        profile = DatasetProfile(
            role=role,
            sheet_name=sheet_name,
            header_row=header_row,
            row_count=row_count,
            column_count=len(frame.columns),
            columns=list(frame.columns),
            inferred_types={column: str(cp.inferred_dtype.value) for column, cp in zip(frame.columns, column_profiles)},
            null_counts=null_counts,
            column_profiles=column_profiles,
        )
        dataset = ParsedDataset(dataframe=frame, profile=profile)
        self._cache[cache_key] = dataset
        return dataset

    def detect_roles(self, file1_path: Path, file2_path: Path) -> RoleDetectionResult:
        g1, p1 = self._header_scores(file1_path)
        g2, p2 = self._header_scores(file2_path)
        score1 = g1 - p1
        score2 = g2 - p2

        # Dual Government-like files
        if (g1 > p1 and g2 > p2) or (g1 > 0 and g2 > 0 and p1 == 0 and p2 == 0):
            return RoleDetectionResult(
                file_1_role=DatasetRole.GOVERNMENT,
                file_2_role=DatasetRole.PURCHASE_REGISTER,
                confidence=0.50,
                is_confident=False,
                reason="Both files appear to be Government GSTR-2B workbooks. Please confirm file role assignment.",
            )

        # Dual Purchase Register-like files
        if (p1 > g1 and p2 > g2) or (p1 > 0 and p2 > 0 and g1 == 0 and g2 == 0):
            return RoleDetectionResult(
                file_1_role=DatasetRole.GOVERNMENT,
                file_2_role=DatasetRole.PURCHASE_REGISTER,
                confidence=0.50,
                is_confident=False,
                reason="Both files appear to be Purchase Register workbooks. Please confirm file role assignment.",
            )

        if score1 > score2 and (g1 > 0 or p2 > 0):
            conf = min(0.98, round(0.70 + (g1 + p2) * 0.05, 2))
            return RoleDetectionResult(
                file_1_role=DatasetRole.GOVERNMENT,
                file_2_role=DatasetRole.PURCHASE_REGISTER,
                confidence=conf,
                is_confident=conf >= 0.80,
                reason=f"File 1 identified as Government GSTR-2B (confidence {conf:.0%}).",
            )
        elif score2 > score1 and (g2 > 0 or p1 > 0):
            conf = min(0.98, round(0.70 + (g2 + p1) * 0.05, 2))
            return RoleDetectionResult(
                file_1_role=DatasetRole.PURCHASE_REGISTER,
                file_2_role=DatasetRole.GOVERNMENT,
                confidence=conf,
                is_confident=conf >= 0.80,
                reason=f"File 2 identified as Government GSTR-2B (confidence {conf:.0%}).",
            )
        else:
            return RoleDetectionResult(
                file_1_role=DatasetRole.GOVERNMENT,
                file_2_role=DatasetRole.PURCHASE_REGISTER,
                confidence=0.50,
                is_confident=False,
                reason="File header semantics are ambiguous. Please confirm file role assignment.",
            )

    @staticmethod
    def _stringify(value: object) -> str:
        if value is None or pd.isna(value):
            return ""
        if isinstance(value, (datetime, date)):
            try:
                return value.isoformat()
            except Exception:
                return str(value)[:80]
        return str(value)[:80]

    def _infer_dtype(self, series: pd.Series) -> tuple[CanonicalDataType, pd.Series | None]:
        try:
            non_null = series.dropna()
            if non_null.empty:
                return CanonicalDataType.STRING, None
            if pd.api.types.is_datetime64_any_dtype(series):
                return CanonicalDataType.DATE, non_null
            if pd.api.types.is_numeric_dtype(series):
                return CanonicalDataType.NUMBER, non_null

            # Fast 50-row strided sample probe across non-null rows
            step = max(1, len(non_null) // 50)
            sample = non_null.iloc[::step][:50].astype(str)

            # Check numeric sample first (fastest)
            try:
                parsed_sample_num = pd.to_numeric(sample, errors="coerce")
                if float(parsed_sample_num.notna().mean()) >= 0.85:
                    parsed_numbers = pd.to_numeric(non_null.astype(str), errors="coerce").dropna()
                    if float(len(parsed_numbers) / len(non_null)) >= 0.95:
                        return CanonicalDataType.NUMBER, parsed_numbers
            except Exception:
                pass

            # Check date sample next
            try:
                parsed_sample_date = pd.to_datetime(sample, errors="coerce", format="mixed", dayfirst=True)
                if float(parsed_sample_date.notna().mean()) >= 0.75:
                    known_formats = (
                        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d-%b-%y", "%Y/%m/%d", "%Y-%m-%dT%H:%M:%S"
                    )
                    str_non_null = non_null.astype(str).str.strip()
                    for fmt in known_formats:
                        try:
                            fast_dates = pd.to_datetime(str_non_null, errors="coerce", format=fmt).dropna()
                            if float(len(fast_dates) / len(non_null)) >= 0.85:
                                return CanonicalDataType.DATE, fast_dates
                        except Exception:
                            continue

                    parsed_dates = pd.to_datetime(str_non_null, errors="coerce", format="mixed", dayfirst=True).dropna()
                    if float(len(parsed_dates) / len(non_null)) >= 0.9:
                        return CanonicalDataType.DATE, parsed_dates
            except Exception:
                pass

            return CanonicalDataType.STRING, None
        except Exception:
            return CanonicalDataType.STRING, None

    def _profile_column(
        self,
        name: str,
        series: pd.Series,
        row_count: int,
        unique_count: int | None = None,
        null_count: int | None = None,
    ) -> ColumnProfile:
        try:
            non_null = series.dropna()
            non_null_count = row_count - null_count if null_count is not None else int(non_null.size)
            inferred, parsed_series = self._infer_dtype(series)
            samples = [self._stringify(value) for value in non_null.head(50).drop_duplicates().head(5).tolist()]
            minimum: str | None = None
            maximum: str | None = None

            if non_null_count and inferred == CanonicalDataType.NUMBER and parsed_series is not None and not parsed_series.empty:
                try:
                    minimum, maximum = self._stringify(parsed_series.min()), self._stringify(parsed_series.max())
                except Exception:
                    pass
            elif non_null_count and inferred == CanonicalDataType.DATE and parsed_series is not None and not parsed_series.empty:
                try:
                    minimum, maximum = self._stringify(parsed_series.min()), self._stringify(parsed_series.max())
                except Exception:
                    pass

            hints: list[str] = []
            if non_null_count:
                try:
                    norm_name = re.sub(r"[^a-z0-9]", "", name.lower())
                    has_gstin_header = any(k in norm_name for k in ("gstin", "ctin", "taxid", "tin", "supplier", "vendor", "party", "counterparty", "seller", "buyer"))
                    if has_gstin_header:
                        hints.append("gstin")
                    else:
                        sample_strings = [str(s).strip() for s in non_null.head(30)]
                        potential_gstin = any(len(s) == 15 and s[:2].isdigit() for s in sample_strings)
                        if potential_gstin:
                            sample_gstin = [s for s in sample_strings if len(s) == 15]
                            if sample_gstin and any(self.GSTIN_PATTERN.match(s.upper()) for s in sample_gstin):
                                string_values = non_null.astype(str).str.strip().str.upper()
                                if float(string_values.str.match(self.GSTIN_PATTERN).mean()) >= 0.5:
                                    hints.append("gstin")
                except Exception:
                    pass

            if inferred == CanonicalDataType.DATE:
                hints.append("date")
            if inferred == CanonicalDataType.NUMBER:
                hints.append("numeric")
                try:
                    if parsed_series is not None and not parsed_series.empty and float(parsed_series.between(0, 100).mean()) >= 0.95:
                        hints.append("percentage_range")
                except Exception:
                    pass

            non_null_percentage = 0.0 if row_count == 0 else round(non_null_count * 100 / row_count, 2)
            u_count = unique_count if unique_count is not None else int(non_null.nunique(dropna=True))
            return ColumnProfile(
                column_name=name,
                inferred_dtype=inferred,
                pandas_dtype=str(series.dtype),
                non_null_count=non_null_count,
                non_null_percentage=non_null_percentage,
                null_percentage=round(100 - non_null_percentage, 2),
                unique_count=u_count,
                sample_values=samples,
                minimum=minimum,
                maximum=maximum,
                pattern_hints=hints,
            )
        except Exception:
            non_null = series.dropna()
            non_null_count = int(non_null.size)
            non_null_percentage = 0.0 if row_count == 0 else round(non_null_count * 100 / row_count, 2)
            samples = [str(v)[:80] for v in non_null.head(5).tolist()]
            return ColumnProfile(
                column_name=name,
                inferred_dtype=CanonicalDataType.STRING,
                pandas_dtype=str(series.dtype),
                non_null_count=non_null_count,
                non_null_percentage=non_null_percentage,
                null_percentage=round(100 - non_null_percentage, 2),
                unique_count=int(non_null.nunique(dropna=True)),
                sample_values=samples,
                minimum=None,
                maximum=None,
                pattern_hints=[],
            )



