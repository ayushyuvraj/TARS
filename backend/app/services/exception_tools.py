from __future__ import annotations

from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path
from time import perf_counter
from uuid import UUID

from openpyxl import load_workbook

from app.domain.models import (
    ActorType, AgentEvent, AmbiguousSelectionRequest, CandidateMatch, CandidateStatus,
    CopilotEvidence, CopilotToolCall, DatasetRole, ExceptionBreakdown, ExceptionRecord,
    ExceptionSearchRequest, ExceptionSearchResult, MatchResult, NearMatchReconciliationSummary,
    PolicySimulationRequest, PolicySimulationResult, ReconciliationPolicy,
    SemanticClassification, VarianceAnalysis,
)
from app.services.near_match import NearMatchEngine
from app.services.reconciliation import ReconciliationNotReadyError, ReconciliationService


class ExceptionToolError(RuntimeError):
    pass


class ExceptionToolService:
    """Controlled deterministic read/action tools used by UI and Copilot orchestration."""

    def __init__(self, reconciliation: ReconciliationService) -> None:
        self.reconciliation = reconciliation

    def _session(self, reconciliation_id: UUID):
        session = self.reconciliation.get(reconciliation_id)
        if session.near_match_analysis is None or session.near_match_summary is None or session.confirmed_mapping is None:
            raise ReconciliationNotReadyError("Completed near-match analysis is required")
        return session

    def _data(self, session):
        mappings = NearMatchEngine._mapping(session.confirmed_mapping)
        government = self.reconciliation.parser.parse(
            Path(session.government_file.stored_path), DatasetRole.GOVERNMENT  # type: ignore[union-attr]
        ).dataframe
        purchase = self.reconciliation.parser.parse(
            Path(session.purchase_register_file.stored_path), DatasetRole.PURCHASE_REGISTER  # type: ignore[union-attr]
        ).dataframe
        return mappings, government, purchase

    def _source_row_values(self, session, role: DatasetRole, record_id: str) -> dict:
        """Read one mapped row without profiling or materializing the workbook."""
        mappings = NearMatchEngine._mapping(session.confirmed_mapping)
        mapping = mappings[role]
        uploaded = session.government_file if role == DatasetRole.GOVERNMENT else session.purchase_register_file
        path = Path(uploaded.stored_path)  # type: ignore[union-attr]
        profile = uploaded.profile  # type: ignore[union-attr]
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[profile.sheet_name]
            header = next(worksheet.iter_rows(
                min_row=profile.header_row, max_row=profile.header_row, values_only=True,
            ))
            indexes = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
            record_column = mapping["record_id"]
            record_index = indexes.get(record_column)
            if record_index is None:
                raise ExceptionToolError(f"Mapped record ID column {record_column} was not found")
            for row in worksheet.iter_rows(min_row=profile.header_row + 1, values_only=True):
                if str(row[record_index]).strip() != record_id:
                    continue
                source_values = {
                    source: row[indexes[source]]
                    for source in mapping.values()
                    if source in indexes
                }
                return self._row_values(source_values, mapping)
        finally:
            workbook.close()
        raise ExceptionToolError(f"Source row for {record_id} was not found")

    @staticmethod
    def _row_values(row, mapping: dict[str, str]) -> dict:
        values = {}
        for canonical, source in mapping.items():
            value = row[source]
            if hasattr(value, "isoformat"):
                value = value.isoformat()
            elif hasattr(value, "item"):
                value = value.item()
            values[canonical] = None if str(value) == "nan" else value
        return values

    def get_reconciliation_summary(self, reconciliation_id: UUID) -> NearMatchReconciliationSummary:
        return self._session(reconciliation_id).near_match_summary

    def get_exception_breakdown(self, reconciliation_id: UUID) -> ExceptionBreakdown:
        session = self._session(reconciliation_id)
        analysis, summary = session.near_match_analysis, session.near_match_summary
        categories = Counter(
            (item.final_category or item.proposed_category).value
            for item in self.reconciliation.repository.list_semantic_classifications(reconciliation_id)
            if item.review_status != "unclassified"
        )
        return ExceptionBreakdown(
            reconciliation_id=reconciliation_id,
            remaining_government=summary.remaining_government_records,
            remaining_purchase_register=summary.remaining_purchase_register_records,
            ambiguous=len(analysis.ambiguities),
            material_mismatch=len(analysis.material_mismatch_government_ids),
            gst_only=len(analysis.gst_only_government_ids),
            pr_only=len(analysis.pr_only_purchase_register_ids),
            semantic_categories=dict(categories),
        )

    def _exception_status(self, analysis, record_id: str) -> tuple[str, DatasetRole] | None:
        if any(item.government_record_id == record_id for item in analysis.ambiguities):
            return "AMBIGUOUS", DatasetRole.GOVERNMENT
        if record_id in analysis.material_mismatch_government_ids:
            return "MATERIAL_MISMATCH", DatasetRole.GOVERNMENT
        if record_id in analysis.gst_only_government_ids:
            return "GST_ONLY", DatasetRole.GOVERNMENT
        if record_id in analysis.pr_only_purchase_register_ids:
            return "PR_ONLY", DatasetRole.PURCHASE_REGISTER
        return None

    def get_record(
        self, reconciliation_id: UUID, record_id: str, *, include_source_values: bool = True,
    ) -> ExceptionRecord:
        session = self._session(reconciliation_id)
        analysis = session.near_match_analysis
        classified = self._exception_status(analysis, record_id)
        if classified is None:
            raise ExceptionToolError(f"Exception record {record_id} was not found")
        status, role = classified
        candidates = self.get_ranked_candidates(reconciliation_id, record_id) if role == DatasetRole.GOVERNMENT else sorted(
            [item for item in analysis.candidates if item.purchase_register_record_id == record_id],
            key=lambda item: (-item.match_score, item.government_record_id),
        )[:10]
        candidate_values = (
            candidates[0].government_values if candidates and role == DatasetRole.GOVERNMENT
            else candidates[0].purchase_register_values if candidates else {}
        )
        values = candidate_values
        if include_source_values or not candidate_values:
            values = self._source_row_values(session, role, record_id)
        semantic = next((item for item in self.reconciliation.repository.list_semantic_classifications(reconciliation_id)
                         if item.record_id == record_id), None)
        return ExceptionRecord(
            record_id=record_id, source_dataset=role, status=status,
            values=values,
            best_candidate=candidates[0] if candidates else None,
            candidate_count=len(candidates), semantic_classification=semantic,
        )

    def get_ranked_candidates(self, reconciliation_id: UUID, record_id: str) -> list[CandidateMatch]:
        analysis = self._session(reconciliation_id).near_match_analysis
        return sorted(
            [item for item in analysis.candidates if item.government_record_id == record_id],
            key=lambda item: (item.rank, -item.match_score, item.purchase_register_record_id),
        )[:10]

    def get_variance_analysis(self, reconciliation_id: UUID, record_id: str) -> VarianceAnalysis:
        # Candidate-backed exception values are already persisted in the near-match
        # analysis. Variance calculation does not need to reopen the source workbook.
        record = self.get_record(reconciliation_id, record_id, include_source_values=False)
        candidate = record.best_candidate
        if candidate is None:
            return VarianceAnalysis(record_id=record_id, status=record.status,
                                    blockers=["No credible candidate was generated within the configured blocking scope"])
        features = candidate.features
        session = self._session(reconciliation_id)
        amount_tolerance = next((float(rule.tolerance.value) for rule in session.confirmed_policy.rules
                                 if rule.canonical_field == "taxable_value" and rule.tolerance), 0)  # type: ignore[union-attr]
        date_tolerance = next((float(rule.tolerance.value) for rule in session.confirmed_policy.rules
                               if rule.canonical_field == "document_date" and rule.tolerance), 0)  # type: ignore[union-attr]
        variances = {
            "taxable_value": float(features.taxable_value_difference),
            "document_date_days": float(features.document_date_difference_days),
            "igst": float(features.igst_difference), "cgst": float(features.cgst_difference),
            "sgst": float(features.sgst_difference), "cess": float(features.cess_difference),
            "invoice_similarity": features.document_number_similarity,
        }
        allowed = {"taxable_value": amount_tolerance, "document_date_days": date_tolerance,
                   "tax_amount": float(analysis_threshold := session.near_match_analysis.thresholds.material_tax_tolerance)}
        blockers = []
        if variances["taxable_value"] > amount_tolerance:
            blockers.append(f"Taxable value variance ₹{variances['taxable_value']:,.2f} exceeds ₹{amount_tolerance:,.2f}")
        for field in ("igst", "cgst", "sgst", "cess"):
            if variances[field] > float(analysis_threshold):
                blockers.append(f"{field.upper()} variance ₹{variances[field]:,.2f} exceeds ₹{float(analysis_threshold):,.2f}")
        if variances["document_date_days"] > date_tolerance:
            blockers.append(f"Document date variance {int(variances['document_date_days'])} days exceeds {int(date_tolerance)} days")
        if record.status == "AMBIGUOUS":
            blockers.append("Multiple similarly strong candidates fall within the configured ambiguity margin")
        return VarianceAnalysis(
            record_id=record_id, candidate_record_id=candidate.purchase_register_record_id,
            status=record.status, variances=variances, allowed_tolerances=allowed, blockers=blockers,
        )

    def get_policy(self, reconciliation_id: UUID) -> ReconciliationPolicy:
        policy = self._session(reconciliation_id).confirmed_policy
        if policy is None:
            raise ReconciliationNotReadyError("Confirmed policy is unavailable")
        return policy

    def _all_exception_records(self, reconciliation_id: UUID) -> list[ExceptionRecord]:
        session = self._session(reconciliation_id)
        analysis = session.near_match_analysis
        mappings, government, purchase = self._data(session)
        frames = {DatasetRole.GOVERNMENT: government, DatasetRole.PURCHASE_REGISTER: purchase}
        row_indexes = {}
        for role, frame in frames.items():
            mapping = mappings[role]
            row_indexes[role] = {
                str(row[mapping["record_id"]]).strip(): row for _, row in frame.iterrows()
            }
        semantic = {item.record_id: item for item in self.reconciliation.repository.list_semantic_classifications(reconciliation_id)}
        candidates = {}
        for item in analysis.candidates:
            candidates.setdefault(item.government_record_id, []).append(item)
        for items in candidates.values():
            items.sort(key=lambda item: (item.rank, -item.match_score, item.purchase_register_record_id))
        status_ids = (
            ("AMBIGUOUS", DatasetRole.GOVERNMENT, [item.government_record_id for item in analysis.ambiguities]),
            ("MATERIAL_MISMATCH", DatasetRole.GOVERNMENT, analysis.material_mismatch_government_ids),
            ("GST_ONLY", DatasetRole.GOVERNMENT, analysis.gst_only_government_ids),
            ("PR_ONLY", DatasetRole.PURCHASE_REGISTER, analysis.pr_only_purchase_register_ids),
        )
        records = []
        for status, role, identifiers in status_ids:
            for record_id in identifiers:
                row = row_indexes[role].get(record_id)
                if row is None:
                    continue
                ranked = candidates.get(record_id, [])[:10]
                records.append(ExceptionRecord(
                    record_id=record_id, source_dataset=role, status=status,
                    values=self._row_values(row, mappings[role]),
                    best_candidate=ranked[0] if ranked else None,
                    candidate_count=len(ranked), semantic_classification=semantic.get(record_id),
                ))
        return records

    def search_records(self, reconciliation_id: UUID, filters: ExceptionSearchRequest) -> ExceptionSearchResult:
        source_filters = any(value is not None for value in (
            filters.vendor, filters.gstin, filters.amount_min, filters.amount_max,
            filters.date_from, filters.date_to, filters.taxable_variance_min,
            filters.taxable_variance_max, filters.match_score_min, filters.match_score_max,
        ))
        if source_filters:
            return self._search_materialized_records(reconciliation_id, filters)

        session = self._session(reconciliation_id)
        analysis = session.near_match_analysis
        semantic = {
            item.record_id: item
            for item in self.reconciliation.repository.list_semantic_classifications(reconciliation_id)
        }
        government_candidates: dict[str, list[CandidateMatch]] = {}
        purchase_candidates: dict[str, list[CandidateMatch]] = {}
        for item in analysis.candidates:
            government_candidates.setdefault(item.government_record_id, []).append(item)
            purchase_candidates.setdefault(item.purchase_register_record_id, []).append(item)
        for candidate_list in (*government_candidates.values(), *purchase_candidates.values()):
            candidate_list.sort(key=lambda item: (item.rank, -item.match_score, item.purchase_register_record_id))
        status_ids = (
            ("AMBIGUOUS", DatasetRole.GOVERNMENT, [item.government_record_id for item in analysis.ambiguities]),
            ("MATERIAL_MISMATCH", DatasetRole.GOVERNMENT, analysis.material_mismatch_government_ids),
            ("GST_ONLY", DatasetRole.GOVERNMENT, analysis.gst_only_government_ids),
            ("PR_ONLY", DatasetRole.PURCHASE_REGISTER, analysis.pr_only_purchase_register_ids),
        )
        records: list[ExceptionRecord] = []
        for status, role, identifiers in status_ids:
            if filters.statuses and status not in filters.statuses:
                continue
            for record_id in identifiers:
                if filters.record_id and filters.record_id.lower() not in record_id.lower():
                    continue
                ranked = (government_candidates if role == DatasetRole.GOVERNMENT else purchase_candidates).get(record_id, [])
                best = ranked[0] if ranked else None
                category = semantic.get(record_id)
                if filters.taxable_variance_min is not None and (
                    best is None or best.features.taxable_value_difference.quantize(Decimal("0.01")) < filters.taxable_variance_min
                ):
                    continue
                if filters.taxable_variance_max is not None and (
                    best is None or best.features.taxable_value_difference.quantize(Decimal("0.01")) > filters.taxable_variance_max
                ):
                    continue
                if filters.match_score_min is not None and (best is None or best.match_score < filters.match_score_min):
                    continue
                if filters.match_score_max is not None and (best is None or best.match_score > filters.match_score_max):
                    continue
                if filters.semantic_category and (category is None or (category.final_category or category.proposed_category) != filters.semantic_category):
                    continue
                if filters.semantic_confidence_min is not None and (category is None or category.confidence < filters.semantic_confidence_min):
                    continue
                if filters.semantic_confidence_max is not None and (category is None or category.confidence > filters.semantic_confidence_max):
                    continue
                if filters.semantic_review_status and (category is None or category.review_status not in filters.semantic_review_status):
                    continue
                if filters.requires_human_review is not None and (category is None or category.requires_human_review != filters.requires_human_review):
                    continue
                source_values = (
                    best.government_values if best and role == DatasetRole.GOVERNMENT
                    else best.purchase_register_values if best else {}
                )
                list_values = {
                    key: source_values[key]
                    for key in ("counterparty_name", "gstin", "taxable_value")
                    if key in source_values
                }
                records.append(ExceptionRecord(
                    record_id=record_id, source_dataset=role, status=status,
                    values=list_values, best_candidate=None,
                    candidate_count=len(ranked), semantic_classification=category,
                ))
        total = len(records)
        return ExceptionSearchResult(
            total=total, offset=filters.offset, limit=filters.limit,
            records=records[filters.offset:filters.offset + filters.limit],
        )

    def _search_materialized_records(
        self, reconciliation_id: UUID, filters: ExceptionSearchRequest,
    ) -> ExceptionSearchResult:
        records = []
        for record in self._all_exception_records(reconciliation_id):
            record_id = record.record_id
            if filters.statuses and record.status not in filters.statuses:
                continue
            if filters.record_id and filters.record_id.lower() not in record_id.lower():
                continue
            values = record.values
            if filters.vendor and filters.vendor.lower() not in str(values.get("counterparty_name", "")).lower():
                continue
            if filters.gstin and filters.gstin.lower() not in str(values.get("gstin", "")).lower():
                continue
            amount = Decimal(str(values.get("taxable_value") or 0))
            if filters.amount_min is not None and amount < filters.amount_min:
                continue
            if filters.amount_max is not None and amount > filters.amount_max:
                continue
            document_date = values.get("document_date")
            try:
                parsed_date = date.fromisoformat(str(document_date)[:10]) if document_date else None
            except ValueError:
                parsed_date = None
            if filters.date_from is not None and (parsed_date is None or parsed_date < filters.date_from):
                continue
            if filters.date_to is not None and (parsed_date is None or parsed_date > filters.date_to):
                continue
            if record.best_candidate:
                variance = record.best_candidate.features.taxable_value_difference.quantize(Decimal("0.01"))
                if filters.taxable_variance_min is not None and variance < filters.taxable_variance_min:
                    continue
                if filters.taxable_variance_max is not None and variance > filters.taxable_variance_max:
                    continue
                if filters.match_score_min is not None and record.best_candidate.match_score < filters.match_score_min:
                    continue
                if filters.match_score_max is not None and record.best_candidate.match_score > filters.match_score_max:
                    continue
            category = record.semantic_classification
            if filters.semantic_category and (category is None or (category.final_category or category.proposed_category) != filters.semantic_category):
                continue
            if filters.semantic_confidence_min is not None and (category is None or category.confidence < filters.semantic_confidence_min):
                continue
            if filters.semantic_confidence_max is not None and (category is None or category.confidence > filters.semantic_confidence_max):
                continue
            if filters.semantic_review_status and (category is None or category.review_status not in filters.semantic_review_status):
                continue
            if filters.requires_human_review is not None and (category is None or category.requires_human_review != filters.requires_human_review):
                continue
            records.append(record)
        total = len(records)
        return ExceptionSearchResult(total=total, offset=filters.offset, limit=filters.limit,
                                     records=records[filters.offset:filters.offset + filters.limit])

    def simulate_policy_change(
        self, reconciliation_id: UUID, record_id: str, hypothetical: PolicySimulationRequest,
    ) -> PolicySimulationResult:
        session = self._session(reconciliation_id)
        variance = self.get_variance_analysis(reconciliation_id, record_id)
        amount_limit = float(hypothetical.taxable_value_tolerance) if hypothetical.taxable_value_tolerance is not None else variance.allowed_tolerances.get("taxable_value", 0)
        date_limit = hypothetical.document_date_tolerance_days if hypothetical.document_date_tolerance_days is not None else variance.allowed_tolerances.get("document_date_days", 0)
        tax_limit = float(hypothetical.tax_amount_tolerance) if hypothetical.tax_amount_tolerance is not None else variance.allowed_tolerances.get("tax_amount", 0)
        blockers = []
        if variance.variances.get("taxable_value", 0) > amount_limit:
            blockers.append("Taxable value variance still exceeds the hypothetical tolerance")
        if variance.variances.get("document_date_days", 0) > date_limit:
            blockers.append("Document date variance still exceeds the hypothetical tolerance")
        for field in ("igst", "cgst", "sgst", "cess"):
            if variance.variances.get(field, 0) > tax_limit:
                blockers.append(f"{field.upper()} variance still exceeds the hypothetical tax tolerance")
        if variance.status == "AMBIGUOUS":
            blockers.append("Candidate ambiguity would still require explicit human selection")
        result = PolicySimulationResult(
            record_id=record_id, current_policy_revision=session.confirmed_policy.revision,  # type: ignore[union-attr]
            hypothetical=hypothetical, would_satisfy=not blockers, blockers=blockers,
        )
        self.reconciliation.repository.add_event(AgentEvent(
            event_type="policy.simulation_completed", reconciliation_id=reconciliation_id,
            actor_type=ActorType.USER, component="exception_tools", result="read_only",
            metadata={"record_id": record_id, "hypothetical": hypothetical.model_dump(mode="json"),
                      "policy_mutated": False, "would_satisfy": result.would_satisfy},
        ))
        return result

    def select_ambiguous_candidate(
        self, reconciliation_id: UUID, government_record_id: str, decision: AmbiguousSelectionRequest,
    ) -> NearMatchReconciliationSummary:
        session = self._session(reconciliation_id)
        analysis, summary = session.near_match_analysis, session.near_match_summary
        ambiguity = next((item for item in analysis.ambiguities if item.government_record_id == government_record_id), None)
        if ambiguity is None:
            raise ExceptionToolError("Record is not an unresolved ambiguous exception")
        if decision.action == "leave_unresolved":
            self.reconciliation.repository.add_event(AgentEvent(
                event_type="ambiguous_candidate.left_unresolved", reconciliation_id=reconciliation_id,
                actor_type=ActorType.USER, component="exception_tools", result="left_unresolved",
                metadata={"government_record_id": government_record_id},
            ))
            return summary
        if not decision.purchase_register_record_id:
            raise ExceptionToolError("A Purchase Register record ID is required for selection")
        candidates = [item for item in analysis.candidates if item.id in ambiguity.candidate_ids]
        selected = next((item for item in candidates if item.purchase_register_record_id == decision.purchase_register_record_id), None)
        if selected is None:
            raise ExceptionToolError("Selected Purchase Register record is not a candidate for this exception")
        consumed = self.reconciliation.repository.list_match_results(reconciliation_id)
        if any(item.purchase_register_record_id == decision.purchase_register_record_id for item in consumed):
            raise ExceptionToolError("Selected Purchase Register record is already consumed")
        match = self.reconciliation._candidate_match(selected, session.confirmed_policy.revision)  # type: ignore[union-attr]
        match.match_type = "human_selected"
        match.rules_satisfied.append("Explicit human selection from ambiguous candidates")
        analysis.ambiguities = [item for item in analysis.ambiguities if item.government_record_id != government_record_id]
        selected.status = CandidateStatus.NEAR_MATCH_APPROVED
        for candidate in candidates:
            if candidate.id != selected.id:
                candidate.status = CandidateStatus.REJECTED_CANDIDATE
        updated = summary.model_copy(update={
            "ambiguous_records": max(0, summary.ambiguous_records - 1),
            "resolved_records": summary.resolved_records + 1,
            "remaining_government_records": summary.remaining_government_records - 1,
            "remaining_purchase_register_records": summary.remaining_purchase_register_records - 1,
        })
        self.reconciliation.repository.save_human_selected_match(reconciliation_id, match, updated, analysis)
        self.reconciliation.repository.add_event(AgentEvent(
            event_type="ambiguous_candidate.human_selected", reconciliation_id=reconciliation_id,
            actor_type=ActorType.USER, component="exception_tools", result="HUMAN_SELECTED",
            approval_status="approved", metadata={"government_record_id": government_record_id,
                                                   "purchase_register_record_id": decision.purchase_register_record_id,
                                                   "candidate_score": selected.match_score},
        ))
        return updated

    @staticmethod
    def traced(tool_name: str, purpose: str, started: float, result_count: int | None = None) -> CopilotToolCall:
        return CopilotToolCall(tool_name=tool_name, purpose=purpose, status="completed",
                               result_count=result_count, duration_ms=round((perf_counter() - started) * 1000, 2))

    @staticmethod
    def evidence(reference_type: str, reference_id: str, facts: dict) -> CopilotEvidence:
        return CopilotEvidence(reference_type=reference_type, reference_id=reference_id, facts=facts)
